import os
import io
import zipfile
import pytest
from fastapi.testclient import TestClient
from backend.main import app
from backend.company_manager import company_manager
from backend.qb_migrator_engine import (
    parse_quickbooks_iif,
    parse_quickbooks_excel,
    parse_quickbooks_qbw,
    parse_quickbooks_zip,
    parse_quickbooks_migration_file,
    execute_quickbooks_migration
)
import openpyxl

client = TestClient(app)

SAMPLE_IIF_CONTENT = """!CLASS\tNAME
CLASS\tSunset Capital Holdings LLC
CLASS\tSunset Capital Holdings LLC:104 Oak St, Austin, TX 78701
CLASS\tSunset Capital Holdings LLC:2908 Depot Road
CLASS\tLone Star Retail Partners LP
CLASS\tLone Star Retail Partners LP:Suite 100 Commercial Plaza

!ACCNT\tNAME\tACCNTTYPE\tDESC\tACCNUM\tOBAL
ACCNT\tOperating Checking Account\tBANK\tPrimary Operating Checking\t10010\t25000.00
ACCNT\tAccounts Receivable\tAR\tUncollected Tenant Rent\t12000\t1500.00
ACCNT\tSecurity Deposits Payable\tOCLIAB\tTenant Escrow Deposits\t22010\t4500.00
ACCNT\tPartner Capital\tEQUITY\tInitial Equity\t30100\t22000.00
ACCNT\tRental Income\tINC\tGross Rent\t40100\t0.00
ACCNT\tRepairs and Maintenance\tEXP\tProperty Repairs\t60100\t0.00
ACCNT\tRepairs and Maintenance:Plumbing\tEXP\tPlumbing Services\t60110\t0.00

!VEND\tNAME\tPRINTAS\tADDR1\tADDR2\tCITY\tSTATE\tZIP\tPHONE1\tEMAIL\tCONT1\tTAXID\tNOTEPAD
VEND\tApex Plumbing LLC\tApex Plumbing\t123 Pipe Way\tSte 200\tAustin\tTX\t78701\t512-555-0101\tbilling@apexplumb.com\tBob Vance\t74-1234567\tMaster Plumber 1099
VEND\tTravis County Tax Office\tTax Office\t5501 Airport Blvd\t\tAustin\tTX\t78751\t512-854-9473\tproperty@traviscountytax.org\tTax Collector\t\tCounty property taxes

!TRNS\tTRNSID\tTRNSTYPE\tDATE\tACCNT\tNAME\tAMOUNT\tDOCNUM\tMEMO
!SPL\tSPLID\tTRNSTYPE\tDATE\tACCNT\tNAME\tAMOUNT\tDOCNUM\tMEMO\tCLASS
!ENDTRNS
TRNS\t\tCHECK\t02/10/2026\tOperating Checking Account\tApex Plumbing LLC\t-550.00\t1001\tEmergency water heater repair
SPL\t\tCHECK\t02/10/2026\tRepairs and Maintenance:Plumbing\tApex Plumbing LLC\t550.00\t1001\tEmergency water heater repair\tSunset Capital Holdings LLC:104 Oak St, Austin, TX 78701
ENDTRNS
"""


def test_parse_quickbooks_iif():
    parsed = parse_quickbooks_iif(SAMPLE_IIF_CONTENT, "Sunset_Holdings_2026.iif")
    assert parsed["status"] == "SUCCESS"
    assert parsed["format"] == "IIF"
    assert "Sunset" in parsed["company_name"]

    # Classes & Properties
    classes = parsed["classes"]
    assert len(classes) == 2
    class_names = [c["name"] for c in classes]
    assert "Sunset Capital Holdings LLC" in class_names
    assert "Lone Star Retail Partners LP" in class_names

    properties = parsed["properties"]
    assert len(properties) == 3
    prop_names = [p["name"] for p in properties]
    assert any("104 Oak St" in p for p in prop_names)
    assert any("2908 Depot Road" in p for p in prop_names)

    # Address parsed for 104 Oak St
    oak_prop = next(p for p in properties if "104 Oak St" in p["name"])
    assert oak_prop["city"] == "Austin"
    assert oak_prop["state"] == "TX"
    assert oak_prop["zip_code"] == "78701"

    # Accounts
    accounts = parsed["accounts"]
    assert len(accounts) == 7
    bank_acct = next(a for a in accounts if a["account_number"] == "10010")
    assert bank_acct["type"] == "ASSET"
    assert bank_acct["balance_total"] == 25000.00

    plumbing_acct = next(a for a in accounts if a["name"] == "Plumbing")
    assert plumbing_acct["is_sub_account"] is True
    assert plumbing_acct["parent_account_name"] == "Repairs and Maintenance"

    # Vendors
    vendors = parsed["vendors"]
    assert len(vendors) == 2
    apex = next(v for v in vendors if "Apex" in v["name"])
    assert apex["tax_id"] == "74-1234567"
    assert apex["is_1099_eligible"] is True
    assert apex["city"] == "Austin"
    assert apex["contact_person"] == "Bob Vance"

    # Transactions
    txns = parsed["transactions"]
    assert len(txns) == 1
    assert txns[0]["payee"] == "Apex Plumbing LLC"
    assert txns[0]["amount"] == 550.00
    assert txns[0]["check_number"] == "1001"
    assert txns[0]["class_name"] == "Sunset Capital Holdings LLC"


def test_parse_quickbooks_excel_multi_sheet():
    wb = openpyxl.Workbook()
    # Sheet 1: Classes
    ws_classes = wb.active
    ws_classes.title = "Class List"
    ws_classes.append(["Class", "Description", "Active"])
    ws_classes.append(["Redwood Capital LLC", "Main Holding Company", "Yes"])
    ws_classes.append(["Redwood Capital LLC:500 Elm St Unit 1", "Elm Street Asset", "Yes"])
    ws_classes.append(["Redwood Capital LLC:500 Elm St Unit 2", "Elm Street Asset", "Yes"])

    # Sheet 2: Chart of Accounts
    ws_coa = wb.create_sheet("Chart of Accounts")
    ws_coa.append(["Account", "Type", "Balance Total", "Description"])
    ws_coa.append(["10010 · Frost Bank Checking", "Bank", "$18,500.00", "Operating Account"])
    ws_coa.append(["40100 · Rental Income", "Income", "$0.00", "Tenant Rents"])
    ws_coa.append(["60100 · Repairs & Maintenance", "Expense", "$0.00", "General Repairs"])

    # Sheet 3: Vendors
    ws_vend = wb.create_sheet("Vendor Contact List")
    ws_vend.append(["Vendor", "Company Name", "Contact", "Phone", "Address", "City", "State", "ZIP", "Tax ID"])
    ws_vend.append(["Reliable Electric", "Reliable Electric Inc", "Tom Edison", "512-555-0144", "789 Shock Rd", "Austin", "TX", "78702", "12-9876543"])

    buf = io.BytesIO()
    wb.save(buf)
    excel_bytes = buf.getvalue()

    parsed = parse_quickbooks_excel(excel_bytes, "Redwood_Portfolio_Export.xlsx")
    assert parsed["status"] == "SUCCESS"
    assert parsed["format"] == "EXCEL"
    assert len(parsed["classes"]) >= 1
    assert len(parsed["properties"]) >= 2
    assert len(parsed["accounts"]) >= 3
    assert len(parsed["vendors"]) >= 1

    vend = parsed["vendors"][0]
    assert vend["name"] == "Reliable Electric"
    assert vend["tax_id"] == "12-9876543"


def test_parse_quickbooks_qbw_encrypted_diagnostics():
    # Synthetic QBW byte payload
    qbw_bytes = b"QUICKBOOKS DESKTOP SYBASE ENCRYPTED PAGE STORE"
    parsed = parse_quickbooks_qbw(qbw_bytes, "MyCompany_2026.qbw")
    assert parsed["format"] == "QBW"
    # When binary does not match valid database page store, handles gracefully
    assert parsed["status"] == "ENCRYPTED_OR_RESTRICTED"
    assert "guidance" in parsed
    assert "Lists to IIF Files" in parsed["guidance"]


def test_parse_quickbooks_zip_archive():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("quickbooks_lists.iif", SAMPLE_IIF_CONTENT)

    zip_bytes = buf.getvalue()
    parsed = parse_quickbooks_zip(zip_bytes, "QuickBooks_Complete_Export.zip")
    assert parsed["status"] == "SUCCESS"
    assert parsed["format"] == "ZIP"
    assert "quickbooks_lists.iif" in parsed["detected_files"]
    assert len(parsed["classes"]) == 2
    assert len(parsed["properties"]) == 3
    assert len(parsed["accounts"]) == 7


def test_parse_quickbooks_migration_file_dispatcher():
    # Test IIF
    iif_res = parse_quickbooks_migration_file(SAMPLE_IIF_CONTENT.encode("utf-8"), "test.iif")
    assert iif_res["format"] == "IIF"

    # Test QBW
    qbw_res = parse_quickbooks_migration_file(b"test qbw", "test.qbw")
    assert qbw_res["format"] == "QBW"

    # Test ZIP
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("export.iif", SAMPLE_IIF_CONTENT)
    zip_res = parse_quickbooks_migration_file(buf.getvalue(), "archive.zip")
    assert zip_res["format"] == "ZIP"


def test_execute_quickbooks_migration_end_to_end():
    parsed = parse_quickbooks_iif(SAMPLE_IIF_CONTENT, "Sunset_Holdings_2026.iif")
    payload = {
        "company_name": "Sunset Test Migration Co",
        "ein": "11-2233445",
        "notes": "Automated migration test run",
        "classes": parsed["classes"],
        "properties": parsed["properties"],
        "accounts": parsed["accounts"],
        "vendors": parsed["vendors"],
        "transactions": parsed["transactions"],
        "create_opening_balances": True,
        "import_transactions": True
    }

    result = execute_quickbooks_migration(payload, company_manager)
    assert result["success"] is True
    assert result["classes_imported"] >= 2
    assert result["properties_imported"] >= 3
    assert result["accounts_imported"] >= 7
    assert result["vendors_imported"] >= 2
    assert result["transactions_imported"] >= 1
    assert "Sunset" in result["company_name"]

    # Verify active company switched
    assert company_manager.active_company_key == result["company_key"]

    # Clean up created company file
    company_manager.delete_company(result["company_key"])


def test_api_quickbooks_migrate_preview_endpoint():
    file_tuple = ("test_lists.iif", io.BytesIO(SAMPLE_IIF_CONTENT.encode("utf-8")), "text/plain")
    response = client.post("/api/quickbooks/migrate/preview", files={"file": file_tuple})
    assert response.status_code == 200
    data = response.json()
    assert data["status"] == "SUCCESS"
    assert data["format"] == "IIF"
    assert len(data["classes"]) == 2
    assert len(data["properties"]) == 3
    assert len(data["accounts"]) == 7
    assert len(data["vendors"]) == 2


def test_api_quickbooks_migrate_convert_endpoint():
    parsed = parse_quickbooks_iif(SAMPLE_IIF_CONTENT, "Sunset_Holdings_2026.iif")
    payload = {
        "company_name": "API Sunset Test Co",
        "ein": "55-6677889",
        "notes": "API migration conversion test",
        "classes": parsed["classes"],
        "properties": parsed["properties"],
        "accounts": parsed["accounts"],
        "vendors": parsed["vendors"],
        "transactions": parsed["transactions"],
        "create_opening_balances": True,
        "import_transactions": True
    }

    response = client.post("/api/quickbooks/migrate/convert", json=payload)
    assert response.status_code == 200
    res_data = response.json()
    assert res_data["success"] is True
    assert res_data["company_name"] == "API Sunset Test Co"
    assert res_data["classes_imported"] >= 2
    assert res_data["properties_imported"] >= 3

    # Cleanup
    company_manager.delete_company(res_data["company_key"])


def test_api_quickbooks_migrate_convert_file_endpoint():
    file_tuple = ("Sunset_Direct_Convert.iif", io.BytesIO(SAMPLE_IIF_CONTENT.encode("utf-8")), "text/plain")
    response = client.post(
        "/api/quickbooks/migrate/convert-file",
        files={"file": file_tuple},
        data={"company_name": "Direct Convert Co", "ein": "99-1122334"}
    )
    assert response.status_code == 200
    res_data = response.json()
    assert res_data["success"] is True
    assert res_data["company_name"] == "Direct Convert Co"
    assert res_data["classes_imported"] >= 2

    # Cleanup
    company_manager.delete_company(res_data["company_key"])
